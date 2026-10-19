import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Styles
title_format = doc.add_heading('AI-Powered Semantic Image Retrieval System', 0)
title_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("Mini-Project Report").alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('1. Project Overview and Role', level=1)
doc.add_paragraph(
    "The AI-Powered Semantic Image Retrieval System is a machine learning mini-project designed to allow users to search through a large collection of images using natural-language text queries (e.g., 'a dog playing in a park' or 'someone wearing a blue shirt'). "
    "Traditional image search engines rely on manual tags and metadata. The role of this project is to build an intelligent, zero-shot search engine that genuinely 'understands' the visual contents of an image using multimodal deep learning."
)

doc.add_heading('2. Workflow of the Project', level=1)
workflow = doc.add_paragraph()
workflow.add_run("1. Dataset Ingestion: ").bold = True
workflow.add_run("A directory of images is imported into the system.\n")
workflow.add_run("2. Preprocessing & Embedding Extraction: ").bold = True
workflow.add_run("Each image is passed through a Vision Transformer to extract high-dimensional semantic embeddings. Concurrently, object detection models and facial recognition models extract structured metadata.\n")
workflow.add_run("3. Vector Indexing: ").bold = True
workflow.add_run("The embeddings are indexed into a FAISS Vector Database for rapid nearest-neighbor search.\n")
workflow.add_run("4. Query Processing: ").bold = True
workflow.add_run("A user inputs a natural language query via the web interface. The text is encoded into the same latent vector space as the images.\n")
workflow.add_run("5. Hybrid Retrieval & Ranking: ").bold = True
workflow.add_run("The system performs a Cosine Similarity vector search, and then adjusts the final scores based on exact object matches and facial identity matches.\n")
workflow.add_run("6. Results Display: ").bold = True
workflow.add_run("The top-K matching images are returned and dynamically rendered in a web dashboard.")

doc.add_heading('3. Key Components and Technologies', level=1)
doc.add_paragraph("The system is modularly designed and heavily relies on modern Python and ML libraries:")
components = doc.add_paragraph()
components.add_run("Backend & API: ").bold = True
components.add_run("FastAPI, Uvicorn, Python 3.13\n")
components.add_run("Database & Storage: ").bold = True
components.add_run("SQLite (for relational metadata), FAISS (for vector similarities)\n")
components.add_run("Frontend: ").bold = True
components.add_run("HTML5, CSS3, JavaScript, Jinja2 Templates\n")
components.add_run("Deep Learning Framework: ").bold = True
components.add_run("PyTorch, Transformers (Hugging Face)\n")
components.add_run("Computer Vision: ").bold = True
components.add_run("OpenCV, Pillow, Ultralytics")

doc.add_heading('4. Models and Algorithms Used', level=1)

doc.add_heading('4.1 OpenAI CLIP (Contrastive Language-Image Pretraining)', level=2)
doc.add_paragraph(
    "Model: clip-vit-base-patch32\n"
    "Role: Used as the core semantic engine. CLIP encodes both the images in the dataset and the user's text queries into a shared 512-dimensional vector space. "
    "Because CLIP was trained on millions of image-text pairs, it intuitively understands abstract concepts, colors, and context without requiring manual labeling."
)

doc.add_heading('4.2 Ultralytics YOLO (You Only Look Once)', level=2)
doc.add_paragraph(
    "Model: YOLOv8 Nano (yolov8n.pt)\n"
    "Role: Used for explicit object detection. While CLIP is great at overall scene context, it can sometimes struggle with counting or specific distinct objects. YOLO scans each image to draw bounding boxes and generate confidence scores for common objects (e.g., car, person, dog, phone)."
)

doc.add_heading('4.3 InsightFace (ArcFace)', level=2)
doc.add_paragraph(
    "Model: buffalo_l\n"
    "Role: Used for facial recognition and personalization. When a user uploads a reference photo of themselves, InsightFace generates a facial embedding map. During searches for queries like 'me at the beach', the system computes face similarity scores to isolate images belonging specifically to the user."
)

doc.add_heading('4.4 FAISS (Facebook AI Similarity Search)', level=2)
doc.add_paragraph(
    "Algorithm: IndexFlatIP (Inner Product)\n"
    "Role: Deep learning models generate floating-point vectors. Comparing a search vector against 3,000 image vectors manually using loops is incredibly slow. FAISS is a highly optimized library written in C++ that performs rapid similarity search and clustering of dense vectors."
)

doc.add_heading('4.5 Cosine Similarity & Hybrid Ranking', level=2)
doc.add_paragraph(
    "Role: The project implements a custom hybrid scoring algorithm. Cosine Similarity is calculated by normalizing the CLIP vectors and computing the dot product. The final ranking score is a weighted combination of:\n"
    "• Semantic Score (CLIP) - 60%\n"
    "• Identity Score (InsightFace) - 25%\n"
    "• Object Score (YOLO) - 15%"
)

doc.save('Project_Report_AI_Image_Retrieval.docx')
print("Document generated successfully.")
